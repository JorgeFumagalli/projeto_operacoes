"""Tratamento auditável do diário. Uso: python tratar_carregamentos_revisado.py arquivo.xlsx pasta_saida

Dependências: pandas e openpyxl. Nenhum dado original é sobrescrito.
"""
from pathlib import Path
from collections import Counter, defaultdict
from datetime import datetime, timedelta
import argparse
import json
import re
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
import pandas as pd
import openpyxl


def clean(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    return v.strip() if isinstance(v, str) else v


def ident(v):
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, (int, float)) and float(v).is_integer():
        return str(int(v))
    return str(v)


def text(v):
    v = clean(v)
    return re.sub(r'\s+', ' ', str(v)).strip() if v is not None else None


def norm(v):
    return (text(v) or '').upper()


def date(v):
    # Correção expressa do responsável: carregamentos 2026 pertencem a 2025.
    if isinstance(v, datetime) and v.year in (2025, 2026):
        return v.replace(year=2025).date().isoformat()
    return None


def number(v):
    v = clean(v)
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    # Texto com separador de milhar isolado é ambíguo: não converter por suposição.
    if re.fullmatch(r'\d{1,3}\.\d{3}', s):
        return None
    try:
        return float(s.replace(',', '.'))
    except ValueError:
        return None


def last_data_row(path, sheet_index):
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    last = 2
    with zipfile.ZipFile(path) as z:
        for _, row in ET.iterparse(z.open(f'xl/worksheets/sheet{sheet_index}.xml'), events=['end']):
            if row.tag != ns + 'row':
                continue
            for c in row:
                col = re.sub(r'\d', '', c.get('r', ''))
                if col in ('B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M'):
                    if c.find(ns+'v') is not None or c.find(ns+'is') is not None:
                        last = max(last, int(row.get('r')))
            row.clear()
    return last


def run(source, out):
    out.mkdir(parents=True, exist_ok=True)
    book = None
    stock_cached = None
    if source.suffix == '.sqlite':
        with sqlite3.connect(source) as con:
            records = pd.read_sql_query('SELECT * FROM originais_carregamentos', con).to_dict('records')
            stock_cached = pd.read_sql_query('SELECT * FROM estoque_snapshot', con)
        for r in records:
            for key,value in list(r.items()):
                if pd.isna(value): r[key] = None
            for c in ('data','lote'):
                value = r.get(c)
                if isinstance(value,str) and re.match(r'^\d{4}-\d{2}-\d{2}',value):
                    try: r[c] = datetime.fromisoformat(value)
                    except ValueError: pass
    else:
        book = openpyxl.load_workbook(source, data_only=True, read_only=True)
        last = last_data_row(source, book.sheetnames.index('Carregamentos') + 1)
        cols = ['codigo_cliente', 'cliente', 'pedido', 'codigo_produto', 'descricao',
                'quantidade', 'lote', 'data', 'motorista', 'placa', 'corte', 'observacoes']
        records = []
        for rownum, row in enumerate(book['Carregamentos'].iter_rows(min_row=3, max_row=last, max_col=13, values_only=True), 3):
            vals = [clean(v) for v in row[1:13]]
            if vals[3] is None: continue
            r = dict(zip(cols, vals)); r['linha_origem'] = rownum; records.append(r)
    originals = pd.DataFrame(records)
    audit = []
    issues = []
    treated = []

    def log(r, campo, old, new, regra):
        audit.append(dict(linha_origem=r['linha_origem'], pedido=r.get('pedido'), campo=campo,
                          original=str(old) if old is not None else None,
                          tratado=str(new) if new is not None else None, regra=regra))

    def issue(r, tipo, detalhe):
        issues.append(dict(linha_origem=r['linha_origem'], pedido=r.get('pedido'), tipo=tipo, detalhe=detalhe))

    for r in records:
        t = {'linha_origem': r['linha_origem']}
        for c in ('codigo_cliente', 'pedido', 'codigo_produto', 'lote'):
            t[c] = ident(r[c])
        for c in ('cliente', 'descricao', 'motorista', 'placa', 'observacoes'):
            t[c] = text(r[c])
        t['quantidade'] = number(r['quantidade'])
        t['quantidade_atendida'] = t['quantidade']
        t['corte_original'] = r['corte']
        t['corte_numero'] = 0.0 if r['corte'] is None else number(r['corte'])
        t['quantidade_faltante'] = t['corte_numero']
        t['corte_status'] = ('zero_por_regra_operacional' if r['corte'] is None else
                             'valor_ambiguo' if t['corte_numero'] is None else
                             'quantidade_faltante_registrada')
        known = (t['quantidade_atendida'] is not None and t['quantidade_atendida'] >= 0
                 and t['quantidade_faltante'] is not None and t['quantidade_faltante'] >= 0)
        t['quantidade_solicitada'] = t['quantidade_atendida'] + t['quantidade_faltante'] if known else None
        t['percentual_atendimento'] = (t['quantidade_atendida'] / t['quantidade_solicitada']
                                      if known and t['quantidade_solicitada'] > 0 else None)
        t['atendimento_status'] = ('calculado' if t['percentual_atendimento'] is not None
                                   else 'indeterminado')
        if r['corte'] is None:
            log(t, 'quantidade_faltante', None, 0, 'Regra confirmada pelo responsável: campo Corte vazio significa ausência de corte.')
        t['data_original'] = str(r['data']) if r['data'] is not None else None
        t['data_tratada'] = date(r['data'])
        corrected_year = isinstance(r['data'],datetime) and r['data'].year == 2026
        t['data_status'] = 'ano_corrigido_usuario' if corrected_year else 'registrada' if t['data_tratada'] else 'ausente_ou_invalida'
        if corrected_year:
            log(t, 'data_tratada', r['data'], t['data_tratada'], 'Correção confirmada pelo responsável em 07/10/2026: ano 2026 para 2025; mês e dia preservados.')
        t['lote_original'] = t.pop('lote')
        t['lote_tratado'] = t['lote_original']
        t['lote_status'] = 'registrado' if t['lote_original'] else 'ausente'
        t['lote_data_ancora'] = None
        t['lote_valor_ancora'] = None
        t['lote_linha_ancora'] = None
        t['pedido_status'] = 'pendente'
        # Apenas embalagens explícitas e sem completar unidades não informadas.
        desc = norm(t['descricao'])
        emb = re.search(r'\b(CX|FD|SC|PCT|UN)\s*', desc)
        t['embalagem'] = emb.group(1) if emb else None
        mass = re.search(r'(?<![\dX])(\d+)\s*[Xx]\s*(\d+(?:[.,]\d+)?)\s*(KG|G)\b', desc)
        t['peso_embalagem_kg'] = (int(mass[1])*float(mass[2].replace(',', '.'))/(1000 if mass[3]=='G' else 1)) if mass else None
        t['peso_status'] = 'extraido_explicito_validar_cadastro' if mass else 'nao_determinado'
        if t['quantidade'] is None or t['quantidade'] <= 0:
            issue(t, 'quantidade_invalida', 'Verificar quantidade original; não imputada.')
        if r['corte'] is not None and t['corte_numero'] is None:
            issue(t, 'corte_ambiguo', 'Quantidade faltante preservada. Confirmar separador decimal/milhar.')
        treated.append(t)

    groups = defaultdict(list)
    for r in treated:
        if r['pedido']:
            groups[r['pedido']].append(r)
        else:
            issue(r, 'pedido_ausente', 'Não preencher número de pedido por proximidade.')
    date_candidates = []
    count_original_multi = 0
    corrected_multi = 0
    unresolved_multi = 0
    for pedido, group in groups.items():
        rawdates = set(r['data_original'] for r in group if r['data_original'])
        count_original_multi += len(rawdates) > 1
        dates = Counter(r['data_tratada'] for r in group if r['data_tratada'])
        clients = set(norm(r['cliente']) for r in group if r['cliente'])
        codes = set(r['codigo_cliente'] for r in group if r['codigo_cliente'])
        client_conflict = len(clients)>1 or len(codes)>1
        winner = None
        if not client_conflict and len(dates)==1:
            winner = next(iter(dates))
        elif not client_conflict and len(dates)>1 and dates.most_common(1)[0][1] > sum(dates.values())/2:
            winner = dates.most_common(1)[0][0]
            corrected_multi += 1
        elif len(dates)>1:
            unresolved_multi += 1
        if winner:
            for r in group:
                if r['data_tratada'] != winner:
                    log(r, 'data_tratada', r['data_original'], winner,
                        'Inferência: data única válida ou maioria estrita no mesmo pedido/cliente; não verificada externamente.')
                    r['data_tratada'] = winner
                    r['data_status'] = 'inferida_pedido'
                r['pedido_status'] = 'data_unica_com_inferencia' if any(x['data_status']=='inferida_pedido' for x in group) else 'data_unica_registrada'
            # Atualizar todas as linhas após finalizar o grupo.
            status = 'data_unica_com_inferencia' if any(x['data_status']=='inferida_pedido' for x in group) else 'data_unica_registrada'
            for r in group:
                r['pedido_status'] = status
        else:
            reason = 'cliente_conflitante' if client_conflict else 'datas_sem_maioria' if len(dates)>1 else 'data_indeterminada'
            for r in group:
                r['pedido_status'] = reason
                issue(r, reason, 'Pedido sem decisão segura; preservado e separado da base de pedidos com data única.')
        if len(rawdates)>1 or client_conflict:
            date_candidates.append(dict(pedido=pedido, linhas=len(group), clientes=len(clients), codigos_cliente=len(codes),
                                        datas_originais=' | '.join(sorted(rawdates)), frequencias_validas=json.dumps(dict(dates)),
                                        data_proposta=winner, status=group[0]['pedido_status']))

    # Cadastro: preencher código só quando nome normalizado possui exatamente um código.
    names = defaultdict(set)
    descs = defaultdict(set)
    for r in treated:
        if r['cliente'] and r['codigo_cliente']:
            names[norm(r['cliente'])].add(r['codigo_cliente'])
        if r['descricao']:
            descs[r['codigo_produto']].add(r['descricao'])
    for r in treated:
        candidates = names[norm(r['cliente'])]
        if not r['codigo_cliente'] and len(candidates)==1:
            r['codigo_cliente'] = next(iter(candidates))
            log(r, 'codigo_cliente', None, r['codigo_cliente'], 'Nome normalizado com um único código registrado no arquivo.')
        if not r['descricao'] and len(descs[r['codigo_produto']])==1:
            r['descricao'] = next(iter(descs[r['codigo_produto']]))
            log(r, 'descricao', None, r['descricao'], 'Código de produto com descrição única registrada.')

    # Âncoras exclusivamente em lotes originais inteiros; nunca encadear estimativas.
    anchors = defaultdict(lambda: defaultdict(list))
    for r in treated:
        lot = r['lote_original']
        if lot and re.fullmatch(r'\d+', lot) and r['data_tratada'] and r['pedido_status'] in ('data_unica_registrada','data_unica_com_inferencia'):
            anchors[r['codigo_produto']][r['data_tratada']].append((int(lot), r['linha_origem']))
    for r in treated:
        if r['lote_original']:
            continue
        d = r['data_tratada']
        if not d or r['pedido_status'] not in ('data_unica_registrada','data_unica_com_inferencia'):
            issue(r, 'lote_nao_estimado', 'Data/pedido pendente; estimativa suspensa.')
            continue
        prior = [a for a in anchors[r['codigo_produto']] if a<=d and a[:4]==d[:4]]
        if not prior:
            issue(r, 'lote_nao_estimado', 'Sem lote original anterior do mesmo produto/ano.')
            continue
        anchor = max(prior)
        values = anchors[r['codigo_produto']][anchor]
        counter = Counter(v for v, _ in values)
        most = counter.most_common(1)[0]
        if len(counter)>1 and most[1] <= sum(counter.values())/2:
            issue(r, 'lote_nao_estimado', 'Lotes na data âncora sem maioria estrita; não escolher arbitrariamente.')
            continue
        lot = most[0]
        days = (datetime.fromisoformat(d)-datetime.fromisoformat(anchor)).days
        estimate = lot + days
        r['lote_tratado'] = str(estimate)
        r['lote_status'] = 'estimado_nao_verificado'
        r['lote_data_ancora'] = anchor
        r['lote_valor_ancora'] = lot
        r['lote_linha_ancora'] = next(line for v, line in values if v==lot)
        log(r, 'lote_tratado', None, estimate,
            f'Estimativa autorizada: lote original {lot} + {days} dias corridos; mesmo produto/ano; âncora {anchor}.')

    # Duplicatas são candidatas; nenhuma é removida automaticamente.
    dup = defaultdict(list)
    for r in treated:
        key = tuple(r.get(c) for c in ('codigo_cliente','cliente','pedido','codigo_produto','descricao','quantidade',
                                      'lote_original','data_original','motorista','placa','corte_original','observacoes'))
        dup[key].append(r)
    for group in dup.values():
        for r in group:
            r['duplicata_candidata'] = len(group)>1
            if len(group)>1:
                issue(r, 'duplicata_candidata', 'Linhas iguais não removidas: confirmar se lançamentos legítimos.')
    df = pd.DataFrame(treated)
    df['apto_pedido_data_unica'] = df['pedido_status'].isin(['data_unica_registrada','data_unica_com_inferencia'])
    df['apto_analise_quantidade'] = df['apto_pedido_data_unica'] & df['quantidade'].gt(0)
    single = df[df['apto_pedido_data_unica']].copy()
    assert single.groupby('pedido')['data_tratada'].nunique().le(1).all()
    assert len(df)==len(originals)
    assert all(r['lote_tratado']==r['lote_original'] for r in treated if r['lote_original'])
    assert not df.data_tratada.dropna().str.startswith('2026').any()

    if stock_cached is not None:
        stock = stock_cached
    else:
        inventory = []
        for rownum, row in enumerate(book['Estoque'].iter_rows(min_row=3,max_row=506,max_col=11,values_only=True),3):
            if row[1] is not None and row[6] is not None:
                inventory.append(dict(linha_origem=rownum, codigo_produto=ident(row[1]), lote_original=ident(row[2]),
                                      descricao=text(row[3]), referencia=text(row[4]), peso_original=text(row[5]),
                                      quantidade=number(row[6]), validade=str(row[7]) if row[7] else None,
                                      local=text(row[8]), venc_dias_original=row[9], observacoes=text(row[10]),
                                      data_snapshot=None, saldo_atualizacao='manual'))
        stock = pd.DataFrame(inventory)
    summary = {
        'linhas_carregamentos': len(df), 'pedidos_identificados': int(df.pedido.nunique()),
        'pedidos_multidata_original_inclui_datas_invalidas':count_original_multi,
        'pedidos_multidata_validas_padronizados_por_maioria':corrected_multi,
        'pedidos_multidata_validas_pendentes':unresolved_multi,
        'datas_inferidas_linhas':int((df.data_status=='inferida_pedido').sum()),
        'datas_ano_2026_corrigido_para_2025':sum(isinstance(r['data'],datetime) and r['data'].year==2026 for r in records),
        'lotes_ausentes_original':int(df.lote_original.isna().sum()),
        'lotes_estimados':int((df.lote_status=='estimado_nao_verificado').sum()),
        'lotes_ainda_ausentes':int(df.lote_tratado.isna().sum()),
        'linhas_base_pedido_data_unica':len(single),
        'linhas_pendentes_pedido_data':int((~df.apto_pedido_data_unica).sum()),
        'linhas_duplicatas_candidatas':int(df.duplicata_candidata.sum()),
        'linhas_estoque':len(stock),
        'codigo_cliente_completado':sum(x['campo']=='codigo_cliente' for x in audit),
        'linhas_com_corte_positivo':int(df.quantidade_faltante.gt(0).sum()),
        'linhas_corte_zero_por_regra_operacional':int((df.corte_status=='zero_por_regra_operacional').sum()),
        'linhas_corte_valor_ambiguo':int((df.corte_status=='valor_ambiguo').sum()),
        'linhas_percentual_atendimento_calculado':int(df.percentual_atendimento.notna().sum()),
    }
    auditdf = pd.DataFrame(audit, columns=['linha_origem','pedido','campo','original','tratado','regra'])
    issuedf = pd.DataFrame(issues, columns=['linha_origem','pedido','tipo','detalhe'])
    rules = pd.DataFrame([
        ('Datas 2026', 'Corrigidas para 2025 por determinação expressa do responsável em 07/10/2026. Apenas carregamentos; originais preservados. Validades não alteradas.'),
        ('Pedido em data única', 'Data única válida ou maioria estrita de linhas com cliente consistente; inferência auditada. Empates/conflitos pendentes.'),
        ('Lote estimado', 'Lote numérico original do mesmo produto na data anterior mais próxima, no mesmo ano, + diferença de dias corridos. Maioria estrita na âncora quando houver vários lotes.'),
        ('Lotes originais', 'Preservados integralmente, inclusive múltiplos lotes por dia, textos ou formatos compostos.'),
        ('Lotes estimados', 'Não representam rastreabilidade comprovada; filtrar lote_status=registrado para medir cobertura real.'),
        ('Corte', 'Confirmado pelo responsável: quantidade que faltou carregar, na unidade do produto. Campo vazio significa zero corte; original preservado e conversão auditada. Separadores ambíguos pendentes.'),
        ('Atendimento do item', 'Qtde confirmado como quantidade atendida. Solicitada=atendida+faltante; percentual=atendida/solicitada. Quantidades atendidas ausentes ou inválidas não são imputadas. Fração de 0 a 1.'),
        ('Cobertura do atendimento', 'Percentuais calculados nas linhas com quantidades válidas, incluindo corte zero pela regra operacional. Não somar unidades distintas para produzir taxa geral. Pedidos/datas pendentes seguem separados.'),
        ('Estoque', 'Saldo atualizado manualmente. Data de referência desconhecida; não calcular vencimento em relação a hoje.'),
        ('Peso/embalagem', 'Extração apenas de padrão explícito NxPesoG/KG. Valores provisórios; não somar caixas/fardos como unidade homogênea.'),
        ('Duplicatas', 'Candidatas sinalizadas e mantidas; não excluir sem comprovação.'),
        ('Dados públicos', 'Pacote de trabalho contém dados identificáveis. Preparar versão anonimizada antes de publicação.'),
    ], columns=['tema','regra'])
    tables = {'Carregamentos_tratados':df,'Base_pedido_data_unica':single,
              'Pendencias':issuedf,'Auditoria':auditdf,'Conferencia_pedidos':pd.DataFrame(date_candidates),
              'Estoque_snapshot':stock,'Originais_carregamentos':originals,
              'Resumo':pd.DataFrame(summary.items(),columns=['indicador','valor']),'Regras':rules}
    # A planilha é montada pelo builder JS separado; este script só trata os dados.
    ordered = ['Resumo','Carregamentos_tratados','Base_pedido_data_unica','Pendencias','Auditoria',
               'Conferencia_pedidos','Estoque_snapshot','Originais_carregamentos','Regras']
    serial = {name: json.loads(tables[name].to_json(orient='split',date_format='iso',force_ascii=False)) for name in ordered}
    (out/'dados_planilha.json').write_text(json.dumps(serial,ensure_ascii=False),encoding='utf-8')
    db=out/'operacoes.sqlite'
    if db.exists():db.unlink()
    with sqlite3.connect(db) as con:
        for sheet,data in tables.items():
            converted=data.copy()
            for c in converted:
                if converted[c].dtype=='object':
                    converted[c]=converted[c].map(lambda x:str(x) if isinstance(x,datetime) else x)
            converted.to_sql(sheet.lower(),con,index=False,if_exists='replace')
        con.execute('CREATE INDEX idx_pedido ON carregamentos_tratados(pedido)')
        con.execute('CREATE INDEX idx_produto_data ON carregamentos_tratados(codigo_produto,data_tratada)')
    (out/'resumo.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    report = '# Base operacional — primeira etapa de tratamento\n\n'
    report += '\n'.join(f'- {k}: {v}' for k,v in summary.items())
    report += '\n\n## Decisões e limites\n\n' + '\n'.join(f'- **{a}:** {b}' for a,b in rules.itertuples(index=False,name=None))
    report += '\n\n## Como utilizar\n\nConsulte operacoes.sqlite e dados_planilha.json. A execução Python não exporta XLSX. Carregamentos_tratados preserva todas as linhas e as colunas originais de lote/data. Base_pedido_data_unica separa registros com data única por pedido; datas inferidas continuam identificadas. Pendencias e Conferencia_pedidos exigem revisão antes de consolidar todos os indicadores. Originais_carregamentos contém os valores em cache do Excel, não suas fórmulas ou layouts. Estoque_snapshot não possui data de referência confirmada.\n\n'
    report += 'Na demonstração pública, os registros são sintéticos. Fontes privadas não devem ser publicadas.\n\n'
    report += 'Reproduzir: `pip install pandas openpyxl` e `python tratar_carregamentos_revisado.py "DIARIO DE CARREGAMENTO 2025.xlsx" saida`.\n'
    (out/'LEIA_ME.md').write_text(report,encoding='utf-8')
    (out/'consultas.sql').write_text('''-- Apenas pedidos com data única; inferências identificadas na origem.
SELECT substr(data_tratada,1,7) AS mes, COUNT(DISTINCT pedido) AS pedidos,
       COUNT(*) AS linhas_itens
FROM base_pedido_data_unica GROUP BY mes ORDER BY mes;

-- Cobertura real de registro de lote, sem contar estimativas como evidência.
SELECT lote_status, COUNT(*) AS linhas FROM carregamentos_tratados GROUP BY lote_status;

-- Quantidade por produto; não somar embalagens diferentes indiscriminadamente.
SELECT codigo_produto, SUM(quantidade) AS quantidade_na_unidade_do_produto
FROM base_pedido_data_unica WHERE quantidade > 0 GROUP BY codigo_produto;

-- Pendências por motivo.
SELECT tipo, COUNT(*) AS ocorrencias FROM pendencias GROUP BY tipo ORDER BY ocorrencias DESC;

-- Quantidades faltantes por produto. Campo Corte vazio equivale a zero, conforme regra confirmada.
SELECT codigo_produto, SUM(quantidade_faltante) AS quantidade_faltante_registrada,
       COUNT(*) AS linhas_com_falta
FROM base_pedido_data_unica WHERE quantidade_faltante > 0 GROUP BY codigo_produto;

-- Atendimento ponderado por produto, em linhas com quantidades válidas e pedido/data consistentes.
-- Não combinar quantidades de produtos em embalagens distintas em uma taxa geral.
SELECT codigo_produto, COUNT(*) AS linhas_com_quantidade_valida,
       SUM(quantidade_atendida) AS atendida, SUM(quantidade_solicitada) AS solicitada,
       1.0 * SUM(quantidade_atendida) / NULLIF(SUM(quantidade_solicitada),0) AS fracao_atendimento
FROM base_pedido_data_unica WHERE atendimento_status = 'calculado'
GROUP BY codigo_produto;
''',encoding='utf-8')
    if book is not None: book.close()
    print(json.dumps(summary,ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('arquivo',type=Path)
    parser.add_argument('saida',type=Path)
    args=parser.parse_args()
    run(args.arquivo,args.saida)
